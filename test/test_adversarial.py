import unittest
import tempfile
from pathlib import Path
from openpyxl import Workbook, load_workbook
from src.services.processors.costar_scoring_preflight import extract_scoring_metrics
from src.services.processors.xlsx_template_filler import XLSXTemplateFiller
from src.services.processors.docling_transformer import DoclingTransformer

class AdversarialEvidenceTests(unittest.TestCase):
    def test_radius_permutations_and_missing_columns(self):
        import itertools
        for radii in itertools.permutations([1,3,5]):
            text='DEMOGRAPHIC SUMMARY\nPopulation '+ ' '.join(f'{r} Mile' for r in radii)+'\n2026 Population '+' '.join(str(r*100) for r in radii)
            self.assertEqual(extract_scoring_metrics([text])['demographics']['population_3mile'],300)
        result=extract_scoring_metrics(['DEMOGRAPHIC SUMMARY\nPopulation 1 Mile 5 Mile\n2026 Population 100 500'])
        self.assertEqual(result['demographics'],{})

    def test_split_pages_signed_growth_and_conflict_are_distinct(self):
        header='DEMOGRAPHIC SUMMARY\nPopulation 1 Mile 3 Mile 5 Mile\n'
        pages=[header+'2026 Population 100 300 500',header+'Median Household Income $10,000 $30,000 $50,000\nPop Growth 2020-2026 -1.0% (2.0)% −3.0%']
        r=extract_scoring_metrics(pages)
        self.assertEqual(r['demographics']['median_hh_income_3mile'],30000)
        self.assertEqual(r['demographics']['population_growth_3mile'],-.02)
        self.assertEqual(r['evidence']['demographics.median_hh_income_3mile'][0]['page'],2)
        r=extract_scoring_metrics(pages+[header+'2026 Population 100 999 500'])
        self.assertNotIn('population_3mile',r['demographics'])
        self.assertIn('demographics.population_3mile',r['conflicts'])

    def test_bad_submarket_text_does_not_erase_other_evidence(self):
        r=extract_scoring_metrics(['DEMOGRAPHIC SUMMARY\nPopulation 1 Mile 3 Mile 5 Mile\n2026 Population 100 300 500','KEY INDICATORS\nBAD HEADER\n12 Mo Delivered Units\nSubmarket unknown'])
        self.assertEqual(r['demographics']['population_3mile'],300)
        self.assertEqual(r['submarket'],{})

    def test_dash_is_unknown_not_zero(self):
        text='KEY INDICATORS\nUnits Vacancy Rate Asking Rent Effective Rent\nSubmarket 1,000 8.0% $1,000 $990 0 0 -\n- 0 8.0% 0.0%\n12 Mo Delivered Units 12 Mo Absorption Units Vacancy Rate 12 Mo Asking Rent Growth'
        r=extract_scoring_metrics([text])['submarket']
        self.assertEqual(r['vacancy_rate'],.08)
        self.assertNotIn('construction_pct_of_inventory',r)
        self.assertNotIn('delivered_pct_of_inventory',r)

    def test_unit_count_comma_and_negative_formats(self):
        transformer=DoclingTransformer()
        for text,value in [('1,234 Units',1234),('-10 Units',-10),('(1,234) Units',-1234),('−45 Units',-45)]:
            self.assertEqual(transformer._parse_units(text),value)

class AdversarialWorkbookTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name)
        self.source=self.root/'template.xlsx';self.output=self.root/'result.xlsx'
        book=Workbook();sheet=book.active;sheet.title='Property';sheet['A1']='Old property';sheet['A2']=999;sheet['B1']='=A2*2';book.save(self.source)
        self.mapping={'json_root':'structured_data[0]','sheets':{'Property':{'mappings':[{'cell':'A1','json_path':'name','type':'text'},{'cell':'A2','json_path':'units','type':'number'}]}}}
    def tearDown(self):self.temp.cleanup()
    def fill(self,data,mapping=None):return XLSXTemplateFiller().fill(str(self.source),{'structured_data':[data]},str(self.output),mapping or self.mapping)
    def test_formula_like_source_is_literal_and_missing_sample_is_cleared(self):
        r=self.fill({'name':'=HYPERLINK("https://example.invalid")'})
        self.assertEqual(r['status'],'success',r)
        book=load_workbook(self.output)
        self.assertEqual(book['Property']['A1'].data_type,'s')
        self.assertIsNone(book['Property']['A2'].value)
        self.assertEqual(book['Property']['B1'].value,'=A2*2')
        self.assertIn('DRAFT',book['Serafina Extraction Review']['B1'].value)
        self.assertEqual(r['readiness'],'draft')
    def test_wrong_sheet_and_nonfinite_values_never_publish(self):
        mapping={'json_root':'structured_data[0]','sheets':{'Missing':{'mappings':[]}}}
        self.assertEqual(self.fill({'name':'New'},mapping)['status'],'error');self.assertFalse(self.output.exists())
        self.assertEqual(self.fill({'name':'New','units':float('inf')})['status'],'error');self.assertFalse(self.output.exists())
    def test_array_overflow_cannot_overwrite_total_or_publish_truncated_rows(self):
        book=load_workbook(self.source);book['Property']['B3']='Total';book.save(self.source)
        mapping={'json_root':'structured_data[0]','sheets':{'Property':{'array_source':'rows','row_start':2,'max_rows':1,'mappings':[{'column':'B','json_field':'name','type':'text'}]}}}
        r=self.fill({'rows':[{'name':'One'},{'name':'Two'}]},mapping)
        self.assertEqual(r['status'],'error',r);self.assertFalse(self.output.exists())
    def test_summary_row_does_not_create_gap_or_shift_formulas(self):
        book=load_workbook(self.source);book['Property']['B4']='Total';book['Property']['C4']='=SUM(C2:C3)';book.save(self.source)
        mapping={'json_root':'structured_data[0]','sheets':{'Property':{'array_source':'rows','row_start':2,'max_rows':2,'mappings':[{'column':'B','json_field':'name','type':'text'}]}}}
        r=self.fill({'rows':[{'bed':'Total','name':'Summary'},{'bed':2,'name':'Actual'}]},mapping)
        self.assertEqual(r['status'],'success',r)
        book=load_workbook(self.output);self.assertEqual(book['Property']['B2'].value,'Actual');self.assertEqual(book['Property']['C4'].value,'=SUM(C2:C3)')

class ProjectionMutationTests(unittest.TestCase):
    def test_metrics_promoted_to_column_headers_do_not_become_unit_mix(self):
        t=DoclingTransformer()
        table={'headers':['ASKING RENTS PER UNIT/SF.Current:','ASKING RENTS PER UNIT/SF.$1,478','ASKING RENTS PER UNIT/SF.$1.71 /SF','VACANCY.Current:','VACANCY.6.0%','VACANCY.11 Units','12 MONTH ABSORPTION.Current:','12 MONTH ABSORPTION.(1,234) Units'], 'rows':[]}
        self.assertEqual(t._classify_tables([table])['unit_breakdown'],[])
        metrics=t._extract_metrics_from_tables([table])
        self.assertEqual(metrics['vacancy']['current']['rate'],6.0)
        self.assertEqual(metrics['absorption']['current'],-1234)

    def test_identical_bed_bath_floorplans_retain_distinct_rows(self):
        t=DoclingTransformer()
        table={'table_id':'p1-t1','headers':['Bed','Bath','Avg SF','Unit Mix.Units'], 'rows':[{'Bed':'1','Bath':'1','Avg SF':'678','Unit Mix.Units':'48'},{'Bed':'1','Bath':'1','Avg SF':'771','Unit Mix.Units':'23'}]}
        rows=t._format_unit_breakdown([table])[0]['rows']
        self.assertEqual(len(rows),2)
        self.assertNotEqual(rows[0]['avgSf'],rows[1]['avgSf'])
        self.assertNotEqual(rows[0]['_source']['row'],rows[1]['_source']['row'])

class CheckpointMutationTests(unittest.TestCase):
    def test_changed_or_truncated_checkpoint_is_recomputed(self):
        from src.services.processors.report_evidence import ReportEvidence
        import json
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp);source=root/'source.pdf';source.write_bytes(b'fixture bytes')
            store=ReportEvidence(source,root,{'batch_pages':8})
            data={'start':1,'end':8,'status':'success','pages':{},'tables':[],'document':{}}
            store.save_batch(data)
            self.assertEqual(store.load_batch(1,8),data)
            path=store.cache/'1-8.json';payload=json.loads(path.read_text());payload['tables']=[{'forged':True}];path.write_text(json.dumps(payload))
            self.assertIsNone(store.load_batch(1,8))
            path.write_text('{interrupted')
            self.assertIsNone(store.load_batch(1,8))

if __name__=='__main__':unittest.main()
