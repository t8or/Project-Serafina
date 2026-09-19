"""Check actual artifacts from verify-local.js against independently labeled fixture values."""
import json
import math
from pathlib import Path
import sys
from openpyxl import load_workbook
verification=json.loads(Path(sys.argv[1]).read_text())
expected={
    'Hawks Landing CoStar.pdf': {'units':144,'vacancy':.097,'income':53216,'inventory':5001,'delivered':9},
    'Serafina CoStart Report.pdf': {'units':183,'vacancy':.06,'income':80544,'inventory':31741,'delivered':1086},
}
for report in verification['reports']:
    values=expected[report['name']]
    workbook=load_workbook(report['workbook']['absolutePath'])
    subject=workbook['Property Summary'];demos=workbook['Demos'];units=workbook['Unit Mix & Comps']
    assert subject['F6'].value==values['units']
    assert math.isclose(subject['F8'].value,values['vacancy'],rel_tol=1e-12)
    assert demos['D8'].value==values['income']
    assert demos['C21'].value==values['inventory']
    assert demos['C22'].value==values['delivered']
    assert demos['C12'].value is None and demos['C13'].value is None
    assert units['E7'].value is None and units['J7'].value is None
    detail_total=sum(cell.value for row in units.iter_rows(min_row=7,max_row=33,min_col=4,max_col=4)
                     for cell in row if isinstance(cell.value,(int,float)))
    assert detail_total==values['units']
    assert workbook['Serafina Extraction Review']['B1'].value.startswith('DRAFT')
    review=workbook['Serafina Extraction Review']
    assert next(row[1] for row in review.values if row[0]=='Source revision')==str(report['revisionId'])
    workbook.close()
    by_property=load_workbook(report['propertyWorkbook']['absolutePath'])
    assert by_property['Demos']['D8'].value==values['income']
    assert by_property['Property Summary']['F6'].value==values['units']
    assert next(row[1] for row in by_property['Serafina Extraction Review'].values if row[0]=='Source revision')==str(report['revisionId'])
    by_property.close()
print('Actual workbook facts, floor-plan totals, and missing-input clearing verified.')
